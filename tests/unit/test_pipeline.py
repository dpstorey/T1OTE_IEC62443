import copy
import json
import subprocess
import sys

import pytest
from conftest import EXAMPLES, SK

import build_plan
import generate_template
import iec_common
import journal
import parse_workbook
import validate_model

MODEL = json.loads((EXAMPLES / "sample-model.json").read_text())


def m(**over):
    d = copy.deepcopy(MODEL)
    d.update(over)
    return d


def codes(rep, kind="errors"):
    return {e["code"] for e in rep[kind]}


# ---- validator ---------------------------------------------------------------


def test_sample_model_is_valid():
    assert validate_model.validate(MODEL)["ok"]


@pytest.mark.parametrize(
    "mutate,code",
    [
        (lambda d: d["members"].append({"zone": "DMZ", "method": "cidr", "value": "10.10.1.0/25"}), "ZONE_OVERLAP"),
        (lambda d: d["members"].append({"zone": "DMZ", "method": "cidr", "value": "10.9.9.5/24"}), "MEMBER_VALUE"),
        (lambda d: d["members"].append({"zone": "DMZ", "method": "cidr", "value": "banana"}), "MEMBER_VALUE"),
        (lambda d: d["members"].append({"zone": "DMZ", "method": "range", "value": "10.2.0.9", "end": "10.2.0.1"}), "MEMBER_VALUE"),
        (lambda d: d["members"].append({"zone": "Nope", "method": "cidr", "value": "10.50.0.0/24"}), "MEMBER_ZONE"),
        (lambda d: d["members"].append({"zone": "DMZ", "method": "magic", "value": "x"}), "MEMBER_METHOD"),
        (lambda d: d["zones"].append({"name": "dmz"}), "ZONE_DUP"),
        (lambda d: d["zones"].append({"name": "Internet"}), "ZONE_RESERVED"),
        (lambda d: d["zones"].append({"name": "Empty"}), "ZONE_EMPTY"),
        (lambda d: d["zones"][0].update(purdue_level=9), "PURDUE_RANGE"),
        (lambda d: d["conduits"][0].update(direction=""), "CONDUIT_DIRECTION"),
        (lambda d: d["conduits"][0].update(zone_b="Ghost"), "CONDUIT_ZONE"),
        (lambda d: d["conduits"][0].update(zone_b="Enterprise"), "CONDUIT_SAME_ZONE"),
        (lambda d: d["conduits"].append({**d["conduits"][0], "name": "again"}), "CONDUIT_DUP_PAIR"),
        (lambda d: d["site"].update(name=""), "SITE_MISSING"),
        (lambda d: d["options"].update(severity="extreme"), "BAD_SEVERITY"),
        (lambda d: d["members"].append({"zone": "DMZ", "method": "filter", "value": "not json"}), "MEMBER_FILTER"),
    ],
)
def test_validator_errors(mutate, code):
    d = copy.deepcopy(MODEL)
    mutate(d)
    rep = validate_model.validate(d)
    assert code in codes(rep) and not rep["ok"]


def test_overlap_downgraded_when_allowed():
    d = copy.deepcopy(MODEL)
    d["members"].append({"zone": "DMZ", "method": "cidr", "value": "10.10.1.0/25"})
    d["options"]["allow_overlap"] = True
    rep = validate_model.validate(d)
    assert rep["ok"] and "ZONE_OVERLAP" in codes(rep, "warnings")


def test_overlap_same_zone_is_fine():
    d = copy.deepcopy(MODEL)
    d["members"].append({"zone": "Control", "method": "cidr", "value": "10.10.1.0/25"})
    assert validate_model.validate(d)["ok"]


def test_warnings():
    d = copy.deepcopy(MODEL)
    d["conduits"][0]["allowed_protocols"] = ["Klingon", "NTP"]
    d["conduits"][1]["allowed_protocols"] = []
    d["zones"].append({"name": "Lonely"})
    d["members"].append({"zone": "Lonely", "method": "vlan", "value": "300"})
    w = codes(validate_model.validate(d), "warnings")
    assert {"PROTOCOL_UNKNOWN", "DIRECTION_SENSITIVE", "CONDUIT_NO_PROTOCOLS", "ZONE_NO_CONDUIT", "NO_OVERLAP_CHECK"} <= w


def test_direction_never_assumed():
    d = copy.deepcopy(MODEL)
    d["conduits"][0]["direction"] = ""
    assert "CONDUIT_DIRECTION" in codes(validate_model.validate(d))


# ---- plan --------------------------------------------------------------------


def _pols(plan):
    return [s for s in plan["steps"] if s["tool"] in build_plan.POLICY_TOOLS.values()]


def _steps(plan, tool):
    return [s for s in plan["steps"] if s["tool"] == tool]


def test_plan_counts_and_state():
    plan = build_plan.build(MODEL)
    pols = _pols(plan)
    assert len(pols) == 14  # 4 boundary-out + 1 unzoned + 4 conformance + 1 remote admin + 4 new asset
    assert all(s["args"]["initial_state"] == "disabled" for s in pols)
    assert plan["summary"]["policies_initial_state"] == "disabled"


def test_plan_without_samples():
    d = copy.deepcopy(MODEL)
    d["options"].update(create_sample_groups=False, create_sample_policies=False)
    pols = _pols(build_plan.build(d))
    assert len(pols) == 4 + 1 + 4


def test_enabled_option_warns():
    d = copy.deepcopy(MODEL)
    d["options"]["policies_enabled"] = True
    plan = build_plan.build(d)
    assert all(s["args"]["initial_state"] == "enabled" for s in _pols(plan))
    assert any("ENABLED" in w for w in plan["warnings"])


def _pol(plan, title_part):
    return next(s for s in _pols(plan) if title_part in s["args"]["title"])["args"]


def test_strict_direction_boundary_sets():
    plan = build_plan.build(MODEL)
    flat = lambda e: [x for clause in e for x in clause]  # noqa: E731
    # DMZ may initiate to Control only; not to Enterprise.
    dmz = flat(_pol(plan, "Boundary | DMZ")["dst"])
    assert any("Control" in x for x in dmz) and not any("Enterprise" in x for x in dmz)
    assert all(x.startswith("!") for x in dmz)
    # Control (bidirectional with Safety) may reach Safety, but not DMZ.
    ctl = flat(_pol(plan, "Boundary | Control")["dst"])
    assert any("Safety" in x for x in ctl) and not any("DMZ" in x for x in ctl)
    # Enterprise reaches DMZ only.
    ent = flat(_pol(plan, "Boundary | Enterprise")["dst"])
    assert any("DMZ" in x for x in ent) and not any("Safety" in x for x in ent)


def test_direction_tolerant_is_symmetric():
    d = copy.deepcopy(MODEL)
    d["options"]["direction_tolerant"] = True
    plan = build_plan.build(d)
    flat = [x for c in _pol(plan, "Boundary | DMZ")["dst"] for x in c]
    assert any("Enterprise" in x for x in flat)


def test_multiblock_zone_is_one_or_clause_and_negations_split():
    plan = build_plan.build(MODEL)
    out = _pol(plan, "Boundary | Control")
    assert len(out["src"]) == 1 and len(out["src"][0]) == 2  # two blocks OR'd
    negs = [c for c in out["dst"] if "Zone | Control" in c[0]]
    assert len(negs) == 2 and all(len(c) == 1 for c in negs)  # one clause per block


def test_conformance_directions_and_negated_protocol_group():
    plan = build_plan.build(MODEL)
    titles = [s["args"]["title"] for s in _pols(plan) if "| Conduit |" in s["args"]["title"]]
    assert len(titles) == 4  # 2 one-way + 1 bidirectional (2)
    one_way = _pol(plan, "Enterprise to DMZ")
    assert one_way["protocol_group_id"].startswith("!$ref:pg:")
    assert not [t for t in titles if "DMZ -> Enterprise" in t]


def test_unzoned_policy_shape():
    plan = build_plan.build(MODEL)
    u = _pol(plan, "unzoned -> zones")
    assert all(len(c) == 1 and c[0].startswith("!") for c in u["src"])
    assert len(u["dst"]) == 1 and not any(x.startswith("!") for x in u["dst"][0])


def test_existing_objects_skipped_and_refs_known():
    name = "IEC62443 | Zone | DMZ"
    existing = {"asset_groups": [{"name": name, "id": "AG9"}],
                "schedule_groups": [{"name": "any time", "id": "S1"}],
                "protocol_groups": [{"name": "Any protocol", "id": "P1"}]}
    plan = build_plan.build(MODEL, existing)
    s = next(s for s in plan["steps"] if s["args"].get("name") == name)
    assert s["action"] == "skip" and s["existing_id"] == "AG9"
    assert plan["known_refs"]["sched:Any Time"] == "S1" and plan["known_refs"]["pg:Any protocol"] == "P1"
    assert not [s for s in plan["steps"] if s["action"] == "lookup"]


def test_asset_names_blocks_plan():
    d = copy.deepcopy(MODEL)
    d["members"].append({"zone": "Safety", "method": "asset_names", "value": "SIS-1, SIS-2"})
    plan = build_plan.build(d)
    assert plan["blocked_on"] and plan["summary"]["needs_input"] == 1


def test_unknown_protocol_left_out_with_warning():
    d = copy.deepcopy(MODEL)
    d["conduits"][0]["allowed_protocols"] = ["Klingon"]
    plan = build_plan.build(d)
    assert any("Klingon" in w for w in plan["warnings"])
    assert not [s for s in _steps(plan, "create_protocol_group") if "Enterprise to DMZ" in s["args"]["name"]]


def test_plan_args_match_create_policy_vocabulary():
    allowed = {"title", "level", "event", "src", "dst", "schedule_group_id", "protocol_group_id", "initial_state"}
    for s in _pols(build_plan.build(MODEL)):
        assert set(s["args"]) <= allowed
        if s["tool"] == "create_conversation_policy":
            assert "event" not in s["args"] and "policy_type" not in s["args"]
        assert s["args"]["level"] in ("low", "medium", "high")


def test_md_renders():
    assert "| # | Phase" in build_plan.render_md(build_plan.build(MODEL))


# ---- journal -----------------------------------------------------------------


def run_j(*a):
    r = subprocess.run([sys.executable, str(SK / "scripts" / "journal.py"), *a], capture_output=True, text=True)
    return r


def test_journal_flow(tmp_path):
    plan = build_plan.build(MODEL)
    pf, jf = tmp_path / "plan.json", tmp_path / "j.jsonl"
    pf.write_text(json.dumps(plan))
    # unresolved ref -> refuses
    pol = _pols(plan)[0]
    r = run_j("resolve", "--plan", str(pf), "--journal", str(jf), "--step", str(pol["n"]))
    assert r.returncode != 0 and "unresolved" in r.stderr
    # record lookups and groups, then it resolves
    for s in plan["steps"]:
        if s["n"] >= pol["n"]:
            break
        run_j("record", "--journal", str(jf), "--step", str(s["n"]), "--tool", s["tool"], "--ref", s["ref"] or "",
              "--status", "ok", "--id", f"ID{s['n']}")
    r = run_j("resolve", "--plan", str(pf), "--journal", str(jf), "--step", str(pol["n"]))
    assert r.returncode == 0, r.stderr
    args = json.loads(r.stdout)["args"]
    assert "$ref" not in json.dumps(args)
    assert args["schedule_group_id"] == "ID1" and args["dst"][0][0].startswith("!ID")
    # undo plan: reverse order, ids filled in
    u = json.loads(run_j("undo-plan", "--plan", str(pf), "--journal", str(jf)).stdout)["undo_calls"]
    assert u and u[0]["undo_of_step"] > u[-1]["undo_of_step"]
    assert all("$id" not in json.dumps(c["args"]) for c in u)


def test_stop_on_error(tmp_path):
    plan = build_plan.build(MODEL)
    pf, jf = tmp_path / "plan.json", tmp_path / "j.jsonl"
    pf.write_text(json.dumps(plan))
    first = plan["steps"][0]["n"]
    run_j("record", "--journal", str(jf), "--step", str(first), "--tool", "x", "--status", "error", "--message", "boom")
    nxt = json.loads(run_j("next", "--plan", str(pf), "--journal", str(jf)).stdout)
    assert nxt["blocked"] is True and nxt["step"] == first
    run_j("record", "--journal", str(jf), "--step", str(first), "--tool", "x", "--status", "skipped")
    nxt = json.loads(run_j("next", "--plan", str(pf), "--journal", str(jf)).stdout)
    assert nxt["blocked"] is False and nxt["step"] != first


# ---- workbook round trip -----------------------------------------------------


def fill(template, out, model):
    from openpyxl import load_workbook

    wb = load_workbook(template)
    row = generate_template.FIRST_DATA_ROW + 1

    def put(ws, rows):
        for i, r in enumerate(rows):
            for c, v in enumerate(r, start=1):
                ws.cell(row=row + i, column=c, value=v)

    wb["Site"].cell(row=row, column=1, value=model["site"]["name"])
    put(wb["Zones"], [[z["name"], z["description"], z["purdue_level"], z["sl_target"], z["criticality"]] for z in model["zones"]])
    put(wb["Address Space"], [[x["zone"], x["method"], x["value"], x.get("end", ""), ""] for x in model["members"]])
    put(wb["Conduits"], [[c["name"], c["zone_a"], c["zone_b"], c["direction"], ", ".join(c["allowed_protocols"]), ""] for c in model["conduits"]])
    wb.save(out)


def test_template_roundtrip(tmp_path):
    t, f = tmp_path / "t.xlsx", tmp_path / "f.xlsx"
    generate_template.build("Demo Plant", str(t))
    empty = parse_workbook.parse(str(t))
    assert not empty["model"]["zones"]  # example rows ignored
    fill(t, f, MODEL)
    got = parse_workbook.parse(str(f))["model"]
    assert [z["name"] for z in got["zones"]] == [z["name"] for z in MODEL["zones"]]
    assert got["conduits"][2]["direction"] == "bidirectional"
    assert got["options"]["policies_enabled"] is False and got["options"]["name_prefix"] == "IEC62443"
    assert validate_model.validate(got)["ok"]
    a = build_plan.build(got)["summary"]
    assert a["policies_create"] == 4 + 1 + 4  # samples off in template defaults


def test_parse_gaps(tmp_path):
    t, f = tmp_path / "t.xlsx", tmp_path / "f.xlsx"
    generate_template.build(None, str(t))
    d = copy.deepcopy(MODEL)
    d["conduits"][0]["direction"] = ""
    d["conduits"][1]["allowed_protocols"] = []
    d["site"]["name"] = ""
    fill(t, f, d)
    gaps = parse_workbook.parse(str(f))["gaps"]
    assert any("Site name" in g for g in gaps)
    assert any("direction is blank" in g for g in gaps)
    assert any("no allowed protocols" in g for g in gaps)


def test_csv_dir(tmp_path):
    (tmp_path / "zones.csv").write_text("Zone name,Description,Purdue level,SL-T,Criticality\nA,x,1,2,high\n")
    (tmp_path / "address_space.csv").write_text("Zone,Method,Value,Range end,Comment\nA,cidr,10.0.0.0/24,,\n")
    (tmp_path / "site.csv").write_text("Site name,Customer,Notes\nS,C,\n")
    got = parse_workbook.parse(str(tmp_path))["model"]
    assert got["zones"][0]["sl_target"] == 2 and got["members"][0]["value"] == "10.0.0.0/24"


def test_protocol_aliases():
    assert iec_common.protocol_key("Modbus/TCP") == "modbus"
    assert iec_common.protocol_key("ENIP") == "ethernet/ip"
    assert iec_common.protocol_key("nope") is None
